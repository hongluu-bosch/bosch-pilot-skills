"""
31service-toolkit - Excel generation utilities.
Generates RID change summary Excel (no longer reads user-edited input).
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("ERROR: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# Change summary columns
CHANGE_COLUMNS = [
    "Product_Type",
    "ARXML_File_Path",
    "Routine_Name",
    "Old_RID_Decimal",
    "Old_RID_Hex",
    "New_RID_Decimal",
    "New_RID_Hex",
    "Status"
]


def generate_change_summary_excel(changes, output_path, command_text=""):
    """
    Generate an Excel summary of RID changes.
    
    Args:
        changes: list of change dicts from rid_calculator.apply_command_to_routines()
        output_path: path to save the Excel file
        command_text: optional command description for a notes sheet
    
    Returns:
        output_path
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Changes"

    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    changed_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    unchanged_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # Headers
    for col_idx, header in enumerate(CHANGE_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data rows
    for row_idx, ch in enumerate(changes, 2):
        routine = ch["routine"]
        ws.cell(row=row_idx, column=1, value=routine.get("product_type", ""))
        ws.cell(row=row_idx, column=2, value=routine.get("arxml_path", ""))
        ws.cell(row=row_idx, column=3, value=routine.get("short_name", ""))
        ws.cell(row=row_idx, column=4, value=ch["old_rid_decimal"])
        ws.cell(row=row_idx, column=5, value=ch["old_rid_hex"])
        ws.cell(row=row_idx, column=6, value=ch["new_rid_decimal"])
        ws.cell(row=row_idx, column=7, value=ch["new_rid_hex"])
        status = "Changed" if ch["changed"] else "Unchanged"
        ws.cell(row=row_idx, column=8, value=status)

        row_fill = changed_fill if ch["changed"] else unchanged_fill
        for col_idx in range(1, len(CHANGE_COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            cell.fill = row_fill

    # Column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 14
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Notes sheet
    if command_text:
        ws_notes = wb.create_sheet(title="Command")
        ws_notes["A1"] = "Command Text"
        ws_notes["A1"].font = Font(bold=True)
        ws_notes["A2"] = command_text
        ws_notes["A3"] = "Changed Count"
        ws_notes["A3"].font = Font(bold=True)
        changed_count = sum(1 for c in changes if c["changed"])
        ws_notes["A4"] = changed_count
        ws_notes["A5"] = "Total Routines"
        ws_notes["A5"].font = Font(bold=True)
        ws_notes["A6"] = len(changes)
        ws_notes.column_dimensions["A"].width = 60

    wb.save(output_path)
    print(f"[Excel] Change summary saved: {output_path}")
    return output_path


if __name__ == "__main__":
    print("generate_excel.py: Use via orchestrator or import generate_change_summary_excel().")
