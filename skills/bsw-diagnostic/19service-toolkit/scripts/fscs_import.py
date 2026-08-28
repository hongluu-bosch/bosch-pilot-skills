"""Round-trip fscs_edit.xlsx back into fscs.json and FSCS_19.txt."""

import json
import pathlib
import re
from typing import Any, Dict, List

from generate_fscs import _render_fscs_19
from io_encoding import save_json, save_text


def run_xlsx_import(workspace: pathlib.Path, args: Any) -> int:
    xlsx_path = workspace / "outputs" / "fscs" / "fscs_edit.xlsx"
    if not xlsx_path.exists():
        print(f"[ERROR] {xlsx_path} not found. Run --phase fscs first.")
        return 1

    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        print("[ERROR] openpyxl not installed")
        return 1

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [str(c.value) for c in ws[1]]

    def idx(name: str) -> int:
        return headers.index(name)

    json_path = workspace / "outputs" / "fscs" / "fscs.json"
    doc = json.loads(json_path.read_text(encoding="utf-8"))

    # Update global snapshot record numbers from the first non-empty row.
    record_numbers: List[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        raw = _str(row[idx("record_number")])
        if raw:
            record_numbers = _parse_record_numbers(raw)
            break
    if record_numbers:
        doc["dtc_snapshot_record_numbers"] = record_numbers

    updated = []
    did_impl_inputs: Dict[str, Dict[str, str]] = {}
    current_did_hex = ""
    current_used = ""
    current_pt = ""
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue

        did_hex_cell = _safe_idx_str(row, idx, "did_hex")
        if did_hex_cell:
            current_did_hex = did_hex_cell
        did_hex = current_did_hex
        if not did_hex:
            continue

        used_cell = _safe_idx_str(row, idx, "used_flag")
        if used_cell:
            current_used = used_cell
        used = current_used.upper() == "Y"

        pt_cell = _safe_idx_str(row, idx, "Product_Type")
        if pt_cell:
            current_pt = pt_cell
        pt = current_pt

        # Collect implementation inputs once per DID; the sheet may have multiple
        # rows per DID because sub-fields expand vertically.
        notes = _safe_idx_str(row, idx, "impl_notes")
        entry = did_impl_inputs.setdefault(did_hex, {"impl_notes": ""})
        if notes and not entry["impl_notes"]:
            entry["impl_notes"] = notes

        for did in doc["dids"]:
            if did["did_hex"] == did_hex:
                did["used"] = used
                did["product_type"] = pt
                _update_sub_field_from_row(did, row, idx)
                updated.append(did_hex)
                break

    # Apply implementation inputs after the row scan so a DID keeps the first
    # non-empty value even when its later sub-field rows are blank.
    for did in doc["dids"]:
        values = did_impl_inputs.get(did["did_hex"])
        if values:
            did["impl_notes"] = values.get("impl_notes", "")

    out_dir = workspace / "outputs" / "fscs"
    save_json(out_dir / "fscs.json", doc)
    save_text(out_dir / "FSCS_19.txt", _render_fscs_19(doc))

    print(f"[OK] xlsx-import complete. Updated {len(updated)} DIDs.")
    print("[AGENT STOP] Choose next step: --phase arxml or --phase doors.")
    return 0


def _str(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _safe_idx_str(row: tuple, idx, name: str) -> str:
    try:
        return _str(row[idx(name)])
    except (ValueError, IndexError):
        return ""


def _parse_record_numbers(value: str) -> List[str]:
    """Parse a string like '01, FF' or '0x01;0x02;0xFF' into canonical hex list."""
    if not value:
        return []
    result: List[str] = []
    seen = set()
    for part in re.split(r"[,;\s]+", value):
        part = part.strip()
        if not part:
            continue
        if part.startswith("$"):
            part = part[1:]
        m = re.match(r"^(0x)?([0-9A-Fa-f]{1,2})$", part)
        if m:
            canonical = f"0x{int(m.group(2), 16):02X}"
            if canonical not in seen:
                seen.add(canonical)
                result.append(canonical)
    return result


def _update_sub_field_from_row(did: Dict[str, Any], row: tuple, idx) -> None:
    """Update the first sub-field of a DID from an xlsx row, if columns exist."""
    sub_fields = did.get("sub_fields", [])
    if not sub_fields:
        return

    sf = sub_fields[0]
    for col_name, key in [
        ("range_min", "range_min"),
        ("range_max", "range_max"),
        ("unit", "unit"),
        ("resolution", "resolution"),
        ("offset", "offset"),
    ]:
        try:
            sf[key] = _str(row[idx(col_name)])
        except (ValueError, IndexError):
            pass
