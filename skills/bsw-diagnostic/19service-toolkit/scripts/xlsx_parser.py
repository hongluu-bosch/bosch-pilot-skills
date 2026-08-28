"""Generic Excel parser for 19-service freeze-frame questionnaires."""

import pathlib
import re
from typing import Any, Dict, List, Optional


def parse_freeze_frame_sheet(xlsx_path: pathlib.Path, product_types: List[str]) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    sheet = _find_sheet(wb)

    header_row = _find_header_row(sheet)
    cols = _resolve_columns(sheet, header_row)

    dids: List[Dict[str, Any]] = []
    current_did: Optional[Dict[str, Any]] = None
    sub_fields: List[Dict[str, Any]] = []

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        if _is_row_strikethrough(sheet, row_idx, cols):
            continue

        did_hex_raw = _cell_str(sheet, row_idx, cols.get("did"))
        did_name_en = _cell_str(sheet, row_idx, cols.get("name_en"))
        length_raw = _cell_str(sheet, row_idx, cols.get("length"))
        did_hex = _canonical_hex(did_hex_raw)

        if did_hex:
            if current_did:
                current_did["sub_fields"] = sub_fields
                dids.append(current_did)
                sub_fields = []
            size = _parse_int(length_raw) or 1
            current_did = {
                "did_hex": did_hex,
                "did_name_en": did_name_en,
                "did_name_zh": _cell_str(sheet, row_idx, cols.get("name_zh")),
                "size_bytes": size,
                "read_fnc": _guess_read_fnc(did_hex, did_name_en),
                "product_type": _infer_product_type(sheet, row_idx, cols, product_types),
                "used": True,
                "sub_fields": []
            }

        sub_fields.append({
            "byte": _cell_str(sheet, row_idx, cols.get("byte")),
            "bit": _cell_str(sheet, row_idx, cols.get("bit")),
            "name_en": _cell_str(sheet, row_idx, cols.get("sub_name_en")),
            "name_zh": _cell_str(sheet, row_idx, cols.get("sub_name_zh")),
            "range_min": _cell_str(sheet, row_idx, cols.get("min")),
            "range_max": _cell_str(sheet, row_idx, cols.get("max")),
            "unit": _cell_str(sheet, row_idx, cols.get("unit")),
            "method_en": _cell_str(sheet, row_idx, cols.get("method_en")),
            "method_zh": _cell_str(sheet, row_idx, cols.get("method_zh")),
        })

    if current_did:
        current_did["sub_fields"] = sub_fields
        dids.append(current_did)

    # Extract snapshot record numbers from the questionnaire; fall back to default.
    record_numbers = _extract_snapshot_record_numbers(sheet, header_row, cols)

    # Parse resolution/offset into structured fields on each sub_field.
    _enrich_sub_fields(dids)

    return {
        "service": "19",
        "subfunction": "0x04",
        "dtc_snapshot_record_numbers": record_numbers,
        "snapshot_record_number_descriptions": {
            "0x01": "The 1st Snapshot record",
            "0x02": "The last Snapshot record",
            "0xFF": "All snapshot record",
        },
        "freeze_frame_class": "DemFreezeFrameClass_RBAPLCUST",
        "freeze_frame_rec_num_class": "DemFreezeFrameRecNumClass_RBAPLCUST",
        "type_of_freeze_frame_record_numeration": "FF_RECNUM_CONFIGURED",
        "product_types": product_types,
        "dids": dids
    }


def _find_sheet(wb) -> Any:
    candidates = []
    for name in wb.sheetnames:
        lower = name.lower()
        if any(k in lower for k in ("$19", "19service", "snapshot", "freezeframe", "冻结帧")):
            candidates.append(name)
    if len(candidates) == 1:
        return wb[candidates[0]]
    if len(candidates) > 1:
        # prefer exact $19
        for c in candidates:
            if "$19" in c:
                return wb[c]
        return wb[candidates[0]]
    return wb[wb.sheetnames[0]]


def _find_header_row(sheet) -> int:
    for row_idx in range(1, min(sheet.max_row, 20) + 1):
        for cell in sheet[row_idx]:
            val = str(cell.value or "").lower().replace("\n", " ")
            if "dtcsnapshotrecordnumber" in val or "did number" in val or "did" in val:
                return row_idx
    return 1


def _resolve_columns(sheet, header_row: int) -> Dict[str, Optional[int]]:
    mapping: Dict[str, List[str]] = {
        "record": ["dtcsnapshotrecordnumber", "dtc snapshot", "dtc extended", "record number"],
        "did": ["did number", "did", "did number(e)", "did编号"],
        "name_en": ["did description(e)", "did description", "did name", "did含义(英文)"],
        "name_zh": ["did description(c)", "did含义(中文)"],
        "length": ["length", "length (bytes)", "长度"],
        "byte": ["byte", "字节"],
        "bit": ["bit", "位"],
        "sub_name_en": ["sub data name(e)", "sub data name", "子信息名称(英文)"],
        "sub_name_zh": ["sub data name(c)", "子信息名称(中文)"],
        "min": ["range,min", "min", "range min"],
        "max": ["range,max", "max", "range max"],
        "unit": ["unit", "单位"],
        "method_en": ["conversion(e)", "conversion", "转换关系(英文)"],
        "method_zh": ["conversion(c)", "转换关系(中文)"],
        "storage": ["storage", "storage pos.", "存储位置"],
        "access": ["access", "访问"],
        "activation": ["activation", "使能"],
        "active_condition": ["active condition", "activecondition", "激活条件"],
    }
    result: Dict[str, Optional[int]] = {k: None for k in mapping}
    for col_idx, cell in enumerate(sheet[header_row], start=1):
        val = str(cell.value or "").lower().replace("\n", " ").strip()
        for key, keywords in mapping.items():
            if result[key] is None and any(kw in val for kw in keywords):
                result[key] = col_idx
    return result


def _is_row_strikethrough(sheet, row_idx: int, cols: Dict[str, Optional[int]]) -> bool:
    for key in ("did", "name_en", "record", "sub_name_en"):
        col = cols.get(key)
        if col:
            cell = sheet.cell(row=row_idx, column=col)
            if cell.font and cell.font.strike:
                return True
    return False


def _cell_str(sheet, row_idx: int, col: Optional[int]) -> str:
    if not col:
        return ""
    val = sheet.cell(row=row_idx, column=col).value
    if val is None:
        return ""
    return str(val).replace("\n", " ").strip()


def _parse_int(value: str) -> Optional[int]:
    m = re.search(r"0x([0-9A-Fa-f]+)", str(value))
    if m:
        return int(m.group(1), 16)
    m = re.search(r"(\d+)", str(value))
    if m:
        return int(m.group(1))
    return None


def _canonical_hex(value: str) -> Optional[str]:
    val = value.strip().lower()
    if not re.match(r"^(0x)?[0-9a-f]+$", val):
        return None
    if not val.startswith("0x"):
        val = "0x" + val
    return f"0x{int(val, 16):04X}"


def _extract_snapshot_record_numbers(sheet, header_row: int, cols: Dict[str, Optional[int]]) -> List[str]:
    """Search the worksheet for snapshot record numbers and normalise them."""
    # Candidate keywords for a cell/label that introduces record numbers.
    keywords = [
        "dtcsnapshotrecordnumber", "snapshot record number", "snapshot record",
        "record number", "冻结帧记录号", "快照记录号",
    ]

    # 1. Prefer an explicit column value from any row below the header.
    record_col = cols.get("record")
    if record_col:
        for row_idx in range(header_row + 1, min(sheet.max_row, header_row + 50) + 1):
            val = _cell_str(sheet, row_idx, record_col)
            if val:
                parsed = _parse_record_numbers(val)
                if parsed:
                    return parsed

    # 2. Search nearby cells for a label followed by a value on the same or next row.
    for row_idx in range(1, min(sheet.max_row, header_row + 10) + 1):
        for col_idx in range(1, min(sheet.max_column, 20) + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            text = str(cell.value or "").lower().replace("\n", " ").strip()
            if any(kw in text for kw in keywords):
                # Look for a number in the same cell after a colon/equal/space
                parsed = _parse_record_numbers(str(cell.value))
                if parsed:
                    return parsed
                # Look in the cell to the right
                if col_idx + 1 <= sheet.max_column:
                    right = _cell_str(sheet, row_idx, col_idx + 1)
                    parsed = _parse_record_numbers(right)
                    if parsed:
                        return parsed
                # Look in the cell below
                if row_idx + 1 <= sheet.max_row:
                    below = _cell_str(sheet, row_idx + 1, col_idx)
                    parsed = _parse_record_numbers(below)
                    if parsed:
                        return parsed

    # 3. Default record numbers used by the toolkit.
    return ["0x01", "0xFF"]


def _parse_record_numbers(value: str) -> List[str]:
    """Parse a string like '01, FF' or '0x01;0x02;0xFF' into canonical hex list."""
    if not value:
        return []
    result: List[str] = []
    seen = set()
    # Split on common separators
    for part in re.split(r"[,;\s]+", value):
        part = part.strip()
        if not part:
            continue
        # Strip leading $ if present
        if part.startswith("$"):
            part = part[1:]
        # Accept hex like 0x01, 01, FF
        m = re.match(r"^(0x)?([0-9A-Fa-f]{1,2})$", part)
        if m:
            canonical = f"0x{int(m.group(2), 16):02X}"
            if canonical not in seen:
                seen.add(canonical)
                result.append(canonical)
    return result


def _enrich_sub_fields(dids: List[Dict[str, Any]]) -> None:
    """Populate structured resolution/offset fields from method_en text."""
    for did in dids:
        for sf in did.get("sub_fields", []):
            method = sf.get("method_en", "")
            sf["resolution"] = _extract_method_part(method, r"[Rr]esolution\s*[=:]\s*([0-9.eE+-]+)")
            sf["offset"] = _extract_method_part(method, r"[Oo]ffset\s*[=:]\s*([0-9.eE+-]+)")


def _extract_method_part(text: str, pattern: str) -> str:
    m = re.search(pattern, text)
    return m.group(1) if m else ""


def _guess_read_fnc(did_hex: str, name_en: str) -> str:
    """Generate a GAC-style read function name from the DID description."""
    # Drop common trailing noise
    clean = re.sub(r"\s+at\s+Last\s+Fault\s+Code\s+Set\s*$", "", name_en, flags=re.IGNORECASE)
    # Title-case each word and remove non-alphanumerics
    words = re.split(r"[^A-Za-z0-9]+", clean)
    suffix = "".join(w[:1].upper() + w[1:] for w in words if w)
    return f"RBAPLCUST_{did_hex.replace('0x', '')}_{suffix}_ReadData"


def _infer_product_type(sheet, row_idx: int, cols: Dict[str, Optional[int]], default_pts: List[str]) -> str:
    if not default_pts:
        return "Common"

    # Use the Active Condition column if it matches a supplied product type
    condition = _cell_str(sheet, row_idx, cols.get("active_condition"))
    if condition:
        lower = condition.lower()
        for pt in default_pts:
            if pt.lower() == lower:
                return pt

    # Fallback: activation column value, if it is an exact product-type match
    activation = _cell_str(sheet, row_idx, cols.get("activation"))
    if activation:
        lower = activation.lower()
        for pt in default_pts:
            if pt.lower() == lower:
                return pt

    return default_pts[0]
