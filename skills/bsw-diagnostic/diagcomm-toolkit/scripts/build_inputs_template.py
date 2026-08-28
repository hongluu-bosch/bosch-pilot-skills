"""Generate the blank ``inputs/DiagComm.xlsx`` template programmatically.

This is a *dev-only* one-shot tool. Run it after touching:
  - ``assets/DiagComm_schema.json``  (new field, tightened range, ...)
  - the standard Bosch CusDiag ``DEFAULT_PATHS`` baked into runtime.py
  - the DOORS skeleton fields (``mode``, ``upload.dry_run``)

Usage (dev / maintainer; ``<skill>`` = the skill checkout, e.g.
``~/.cursor/skills/diagcomm-toolkit``)::

    python <skill>/scripts/build_inputs_template.py
    python <skill>/scripts/build_inputs_template.py --output path/to/file.xlsx

Output (default): ``assets/inputs_template.xlsx``. The Excel-driven
loader reads from ``inputs/DiagComm.xlsx``; the workflow is
"copy template -> edit". Both files share the exact same shape so the
template doubles as the validating round-trip fixture.

Why programmatic instead of hand-painted xlsx
---------------------------------------------
The 21 parameters + their type / range / enum allow-list change
faster than a manually maintained xlsx can keep up. Sourcing the
template from the schema makes ``assets/DiagComm_schema.json`` the
*single source of truth* and prevents Excel <-> schema drift.

Sheet layout
------------
1. ``Project & Parameters``     -- project identity + 21 diag params
2. ``Paths & Options``          -- file paths + runtime flags
3. ``DOORS Upload``             -- DOORS module UUID + upload mode
4. ``README``                   -- inline how-to-fill guide

Every data sheet uses the canonical 3-column layout:

    A: Field   -- dotted path key (e.g. ``project.name``); the loader
                  matches on this column verbatim
    B: Value   -- editable cell with DataValidation matching the
                  schema (enum dropdown, range whole/decimal, etc.)
    C: Notes   -- read-only hint (default, unit, allowed values,
                  short description)

Required fields get a red fill so the user sees at a glance what
must be touched. Optional fields ship pre-filled with the schema
default; the user can override or leave them.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

# Lazy import so a missing openpyxl gives a clean error message.
try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # noqa: BLE001
    sys.stderr.write(
        "[build_inputs_template] openpyxl is required. "
        "Install with: pip install openpyxl\n"
    )
    raise SystemExit(2) from exc


SCRIPTS_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPTS_DIR.parent
ASSETS_DIR = SKILL_ROOT / "assets"
SCHEMA_PATH = ASSETS_DIR / "DiagComm_schema.json"

DEFAULT_OUTPUT = ASSETS_DIR / "inputs_template.xlsx"

# ---------------------------------------------------------------------------
# Defaults that mirror runtime.DEFAULT_PATHS / DEFAULT_OPTIONS. Duplicated
# here (rather than imported from runtime.py) to keep this script standalone
# and free of pipeline imports.

DEFAULT_PATHS: dict[str, str] = {
    # v2.0.0: workspace lives at <project>/.DCOM_AI/DiagComm_Toolkit_PRJ/,
    # so two ".." segments land back at the project root (was ../../..
    # in v1 when the workspace WAS the skill checkout three levels deep).
    "base_dir": "../..",
    "dcom_root": "*/rb/as/*/core/app/dcom",
    "cantp_common": "RBAPLCust/cfg/Common/CanTp_CusDiag_EcucValues.arxml",
    "cantp_feature_file": "Cubas/cfg/CanTp_Feature_EcucValues.arxml",
    "dcm_common": "RBAPLCust/cfg/Common/Dcm_CusDiag_Can_EcucValues.arxml",
    "dcm_feature_file": "Cubas/cfg/Dcm_Feature_EcucValues.arxml",
    "dcm_services_common": "RBAPLCust/cfg/Common/Dcm_CusDiag_Services_EcucValues_SingleCANID.arxml",
    "can_pt_file": "RBAPLCust/cfg/{product_type}/Can{can_channel}_CusDiag_EcucValues_{product_type}.arxml",
}

DEFAULT_OPTIONS: dict[str, Any] = {
    "dry_run_default": True,
    "validate_before_apply": True,
}

# ---------------------------------------------------------------------------
# Cell styling primitives

HDR_FILL  = PatternFill("solid", fgColor="305496")
HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
NOTES_FILL = PatternFill("solid", fgColor="EDEDED")
REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, italic=True, size=10, color="1F3864")

THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")

# Sentinel used in Field column to denote a section divider row (no data).
SECTION_PREFIX = "## "


# ---------------------------------------------------------------------------
# Sheet builders

def _set_widths(ws, widths: list[int]) -> None:
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w


def _write_header(ws, row: int = 1) -> None:
    for col, text in enumerate(("Field", "Value", "Notes"), start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _write_section(ws, row: int, label: str) -> int:
    """Write a section divider row spanning all 3 columns. Returns ``row+1``."""
    cell = ws.cell(row=row, column=1, value=f"{SECTION_PREFIX}{label}")
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    return row + 1


def _write_data_row(
    ws,
    row: int,
    field: str,
    value: Any,
    notes: str,
    *,
    required: bool = False,
) -> None:
    f = ws.cell(row=row, column=1, value=field)
    v = ws.cell(row=row, column=2, value=value)
    n = ws.cell(row=row, column=3, value=notes)

    f.font = Font(name="Consolas", size=10)
    f.alignment = Alignment(vertical="top")
    f.border = BORDER

    v.alignment = Alignment(vertical="top", wrap_text=False)
    v.border = BORDER
    if required and (value is None or value == ""):
        v.fill = REQUIRED_FILL

    n.fill = NOTES_FILL
    n.font = Font(size=9, color="595959")
    n.alignment = WRAP
    n.border = BORDER


def _attach_validation(ws, dv: DataValidation, target_row: int) -> None:
    """Attach a data validation to the Value cell of a given row."""
    cell_ref = f"B{target_row}:B{target_row}"
    dv.add(cell_ref)
    if dv not in ws.data_validations.dataValidation:
        ws.add_data_validation(dv)


def _make_list_dv(values: list[Any], *, allow_blank: bool = True) -> DataValidation:
    quoted = '"' + ",".join(str(v) for v in values) + '"'
    return DataValidation(
        type="list",
        formula1=quoted,
        allow_blank=allow_blank,
        showDropDown=False,
        errorTitle="Invalid value",
        error=f"Pick one of: {', '.join(map(str, values))}",
    )


def _make_int_range_dv(lo: int, hi: int) -> DataValidation:
    return DataValidation(
        type="whole",
        operator="between",
        formula1=str(lo),
        formula2=str(hi),
        allow_blank=True,
        errorTitle="Out of range",
        error=f"Integer between {lo} and {hi}",
    )


def _make_decimal_min_dv(lo: float) -> DataValidation:
    return DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1=str(lo),
        allow_blank=True,
        errorTitle="Out of range",
        error=f"Number >= {lo}",
    )


def _format_notes(spec: dict[str, Any]) -> str:
    parts: list[str] = []
    desc = spec.get("description") or ""
    if desc:
        parts.append(desc.strip())

    meta_bits: list[str] = []
    if spec.get("type"):
        meta_bits.append(f"type={spec['type']}")
    if "default" in spec:
        meta_bits.append(f"default={spec['default']!r}")
    if isinstance(spec.get("allowed"), list):
        meta_bits.append("allowed=" + " / ".join(map(str, spec["allowed"])))
    rng = []
    if "min" in spec:
        rng.append(str(spec["min"]))
    if "max" in spec:
        rng.append(str(spec["max"]))
    if rng:
        meta_bits.append("range=" + "..".join(rng))
    if spec.get("unit"):
        meta_bits.append(f"unit={spec['unit']}")
    if spec.get("prompt_required"):
        meta_bits.append("REQUIRED")
    if spec.get("maps_to"):
        meta_bits.append(f"-> {spec['maps_to']}")
    if meta_bits:
        parts.append("  [" + ", ".join(meta_bits) + "]")
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Sheet 1: Project & Parameters

def _build_sheet_project(wb: Workbook, schema: dict[str, Any]) -> None:
    ws = wb.create_sheet("Project & Parameters")
    _set_widths(ws, [36, 22, 80])
    _write_header(ws)
    row = 2

    fields = schema.get("fields") or {}

    # --- Section: project identity -----------------------------------------
    row = _write_section(ws, row, "Project identity (REQUIRED)")

    _write_data_row(
        ws, row, "project.name", "",
        "Free text. Used in FSCS header + DOORS Object Text.\n"
        "  [type=string, REQUIRED]",
        required=True,
    )
    row += 1

    pt_spec = fields.get("product_type") or {}
    _write_data_row(
        ws, row, "project.product_type", "",
        _format_notes(pt_spec),
        required=True,
    )
    pt_allowed = pt_spec.get("allowed") or []
    if pt_allowed:
        _attach_validation(ws, _make_list_dv(pt_allowed), row)
    row += 1

    # --- Section: CAN basics ----------------------------------------------
    row = _write_section(ws, row, "CAN basics (defaults usually OK)")
    for key in ("CAN_Channel", "CAN_ID_Format", "Addressing_Method"):
        spec = fields.get(key) or {}
        default = spec.get("default", "")
        _write_data_row(
            ws, row, f"parameters.{key}", default,
            _format_notes(spec),
            required=bool(spec.get("prompt_required")),
        )
        if isinstance(spec.get("allowed"), list):
            _attach_validation(ws, _make_list_dv(spec["allowed"]), row)
        elif spec.get("type") == "int":
            _attach_validation(ws, _make_int_range_dv(0, 31), row)
        row += 1

    # --- Section: CAN_DLC sub-fields --------------------------------------
    row = _write_section(ws, row, "CAN DLC (REQUIRED -- per-direction frame type & length)")
    can_dlc_spec = (fields.get("CAN_DLC") or {}).get("fields") or {}
    for sub_key, sub_spec in can_dlc_spec.items():
        _write_data_row(
            ws, row, f"parameters.CAN_DLC.{sub_key}", "",
            _format_notes(sub_spec),
            required=bool(sub_spec.get("prompt_required")),
        )
        if isinstance(sub_spec.get("allowed"), list):
            _attach_validation(ws, _make_list_dv(sub_spec["allowed"]), row)
        row += 1

    # --- Section: CAN IDs -------------------------------------------------
    row = _write_section(ws, row, "CAN IDs (REQUIRED -- hex with 0x prefix)")
    for key in ("CAN_Functional_Request_ID",
                "CAN_Physical_Request_ID",
                "CAN_Response_ID"):
        spec = fields.get(key) or {}
        _write_data_row(
            ws, row, f"parameters.{key}", "",
            _format_notes(spec),
            required=bool(spec.get("prompt_required")),
        )
        row += 1

    # --- Section: Network-layer timing ------------------------------------
    row = _write_section(ws, row, "Network-layer timing (project-level overrides common)")
    for key in ("N_As", "N_Ar", "N_Bs", "N_Br", "N_Cs", "N_Cr",
                "BS", "STmin"):
        spec = fields.get(key) or {}
        default = spec.get("default", "")
        _write_data_row(
            ws, row, f"parameters.{key}", default,
            _format_notes(spec),
            required=bool(spec.get("prompt_required")),
        )
        if "min" in spec and "max" in spec and spec.get("type") == "float":
            ws.add_data_validation(DataValidation(
                type="decimal", operator="between",
                formula1=str(spec["min"]), formula2=str(spec["max"]),
                allow_blank=True,
                errorTitle="Out of range",
                error=f"Decimal between {spec['min']} and {spec['max']}",
            ))
            ws.data_validations.dataValidation[-1].add(f"B{row}")
        elif spec.get("type") == "float":
            _attach_validation(ws, _make_decimal_min_dv(0.0), row)
        elif spec.get("type") == "int":
            _attach_validation(ws, _make_int_range_dv(0, 65535), row)
        row += 1

    # --- Section: P2 timer ------------------------------------------------
    row = _write_section(ws, row, "Application-layer (P2) timer")
    for key in ("P2_Max", "P2_Star_Max"):
        spec = fields.get(key) or {}
        _write_data_row(
            ws, row, f"parameters.{key}", spec.get("default", ""),
            _format_notes(spec),
        )
        _attach_validation(ws, _make_decimal_min_dv(0.0), row)
        row += 1

    # --- Section: misc ----------------------------------------------------
    row = _write_section(ws, row, "Misc (defaults usually OK)")
    for key in ("PaddingByte", "StrictDlcCheck", "NRC78_Times"):
        spec = fields.get(key) or {}
        _write_data_row(
            ws, row, f"parameters.{key}", spec.get("default", ""),
            _format_notes(spec),
        )
        if spec.get("type") == "bool":
            _attach_validation(ws, _make_list_dv(["TRUE", "FALSE"]), row)
        elif spec.get("type") == "int" and "min" in spec and "max" in spec:
            _attach_validation(ws, _make_int_range_dv(spec["min"], spec["max"]), row)
        row += 1


# ---------------------------------------------------------------------------
# Sheet 2: Paths & Options

PATH_NOTES: dict[str, str] = {
    "base_dir": (
        "Workspace-relative anchor for everything below. Default `../..` "
        "lands on the project root (workspace lives at "
        "<project>/.DCOM_AI/DiagComm_Toolkit_PRJ/, two levels deep)."
    ),
    "dcom_root": (
        "Glob-style relative path from base_dir to the dcom tree. "
        "Wildcards allowed (`*` / `?`). Resolution prefers directories "
        "containing RBAPLCust/ + Cubas/."
    ),
    "cantp_common": "Relative to dcom_root.",
    "cantp_feature_file": "Relative to dcom_root.",
    "dcm_common": "Relative to dcom_root.",
    "dcm_feature_file": "Relative to dcom_root.",
    "dcm_services_common": "Relative to dcom_root.",
    "can_pt_file": (
        "Per-product CAN PT arxml. `{product_type}` and `{can_channel}` "
        "are filled at runtime from project.product_type / parameters.CAN_Channel."
    ),
}

OPTION_NOTES: dict[str, str] = {
    "dry_run_default": (
        "When TRUE, `pipeline.py apply` defaults to dry-run (no file "
        "writes). Override per-run with --apply."
    ),
    "validate_before_apply": (
        "When TRUE, `pipeline.py apply` runs `validate` first and "
        "refuses to apply if any field is out of range."
    ),
}


def _build_sheet_paths(wb: Workbook) -> None:
    ws = wb.create_sheet("Paths & Options")
    _set_widths(ws, [36, 60, 60])
    _write_header(ws)
    row = 2

    row = _write_section(ws, row, "Paths (standard Bosch CusDiag layout; rarely edit)")
    for key, default in DEFAULT_PATHS.items():
        notes = PATH_NOTES.get(key, "")
        _write_data_row(ws, row, f"paths.{key}", default,
                        notes + f"\n  [default={default!r}]")
        row += 1

    row = _write_section(ws, row, "Runtime options")
    for key, default in DEFAULT_OPTIONS.items():
        notes = OPTION_NOTES.get(key, "")
        # bools rendered as TRUE/FALSE so the dropdown matches
        rendered = "TRUE" if default is True else ("FALSE" if default is False else default)
        _write_data_row(ws, row, f"options.{key}", rendered,
                        notes + f"\n  [type=bool, default={default}]")
        _attach_validation(ws, _make_list_dv(["TRUE", "FALSE"]), row)
        row += 1


# ---------------------------------------------------------------------------
# Sheet 3: DOORS Upload

def _build_sheet_doors(wb: Workbook) -> None:
    ws = wb.create_sheet("DOORS Upload")
    _set_widths(ws, [28, 50, 80])
    _write_header(ws)
    row = 2

    row = _write_section(ws, row, "DOORS module identity (REQUIRED for Step 8)")
    _write_data_row(
        ws, row, "doors.document_uuid", "PUT-DOORS-DOCUMENT-UUID-HERE",
        "The fixed UUID of the DOORS module to read/write. Look it up "
        "in DOORS (Properties -> URL/UUID) and paste verbatim. "
        "doors_sync refuses to upload while this is the placeholder.\n"
        "  [type=string, REQUIRED]",
        required=True,
    )
    row += 1

    row = _write_section(ws, row, "Upload behaviour (defaults usually OK)")
    _write_data_row(
        ws, row, "mode", "insert",
        "How rows land in DOORS. `insert` (default) places new rows "
        "after the anchor heading. `update` overwrites an existing "
        "row by AbsoluteNumber (use --update-abs <arxml>=<n> on CLI).\n"
        "  [allowed=insert / update, default='insert']",
    )
    _attach_validation(ws, _make_list_dv(["insert", "update"]), row)
    row += 1

    _write_data_row(
        ws, row, "upload.dry_run", "FALSE",
        "When TRUE, doors_sync builds the upload xlsx but does NOT call "
        "the DOORS MCP server. Useful for inspecting outputs/doors_upload.xlsx.\n"
        "  [type=bool, default=FALSE]",
    )
    _attach_validation(ws, _make_list_dv(["TRUE", "FALSE"]), row)
    row += 1


# ---------------------------------------------------------------------------
# Sheet 4: README

README_LINES: list[str] = [
    "diagcomm-toolkit -- inputs/DiagComm.xlsx",
    "",
    "This file is your ONLY editable input. Every other config the skill needs ",
    "is bundled in assets/ (read-only) and merged with what you put here.",
    "",
    "------------------------------------------------------------------------",
    "How to fill it",
    "------------------------------------------------------------------------",
    "",
    "1. Sheet 'Project & Parameters'",
    "   - Cells with red background are REQUIRED. Fill them or the pipeline ",
    "     refuses to run with a precise list of what's missing.",
    "   - 'project.name' is free text (e.g. 'MyProject_Y2026').",
    "   - Other parameters that ship pre-filled with defaults are project-",
    "     spec sensitive. Compare against your project's diagnostic spec ",
    "     before running 'apply'; override anything that doesn't match.",
    "",
    "2. Sheet 'Paths & Options'",
    "   - Standard Bosch CusDiag layout is pre-filled. ~99% of projects ",
    "     don't touch this sheet at all.",
    "   - Edit only when your arxml tree lives in a non-standard location.",
    "",
    "3. Sheet 'DOORS Upload'  (only needed for Step 8)",
    "   - Replace 'PUT-DOORS-DOCUMENT-UUID-HERE' with the real DOORS module ",
    "     UUID before the first 'doors_sync' run.",
    "   - 'mode' defaults to 'insert'; switch to 'update' only when you ",
    "     intentionally overwrite an existing row by AbsoluteNumber.",
    "",
    "------------------------------------------------------------------------",
    "Behind the scenes",
    "------------------------------------------------------------------------",
    "",
    " - scripts/excel_loader.py reads this xlsx and writes .cache/inputs.json ",
    "   (gitignored) which the rest of the pipeline consumes.",
    " - The DOORS column mapping skeleton (anchor/value_maps/columns/...) ",
    "   lives in assets/doors_mapping_skeleton.yaml and is merged with the ",
    "   3 fields from the 'DOORS Upload' sheet at upload time.",
    " - Every run that needs your inputs auto-regenerates the cache; you ",
    "   never call the loader manually.",
    "",
    "Cell colours",
    " - Blue header     : column titles (do not delete row 1).",
    " - Light-blue band : section divider (do not delete; loader skips them).",
    " - Light-pink Value: REQUIRED field, currently empty.",
    " - Grey Notes col  : read-only hint; safe to ignore in scripts.",
    "",
    "Validation",
    " - Enum cells (product_type, frame_type, mode, ...) carry an Excel ",
    "   data-validation dropdown. Type-anything values are rejected at ",
    "   load time by scripts/excel_loader.py.",
    "",
    "Trouble-shooting",
    " - Lost the file? Reset from the blank template (DESTRUCTIVE — overwrites",
    "   any existing workbook):",
    "       cd <project-root>",
    "       python <skill>/scripts/pipeline.py --init-project --force",
    " - 'invalid xlsx' on load? See <skill>/reference/excel_format.md for the",
    "   schema the loader expects.",
    " - Old user with inputs/*.json from <= 1.19.x? Run:",
    "       cd <project-root>",
    "       python <skill>/scripts/pipeline.py --init-project",
    "       python <skill>/scripts/migrate_v1_19_to_xlsx.py --legacy-from <dir-with-old-json>",
]


def _build_sheet_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 100
    for idx, line in enumerate(README_LINES, start=1):
        cell = ws.cell(row=idx, column=1, value=line)
        cell.alignment = Alignment(vertical="top", wrap_text=False)
        if idx == 1:
            cell.font = Font(bold=True, size=14)
        elif line.startswith("--"):
            cell.font = Font(italic=True, color="595959")
        elif line and line[0].isalpha() and line.endswith(":"):
            # subsection mini-header
            cell.font = Font(bold=True)
        else:
            cell.font = Font(name="Consolas", size=10)
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Top-level entry

def build_workbook(schema: dict[str, Any]) -> Workbook:
    wb = Workbook()
    # Drop the default first sheet; we add our own four below in order.
    default_ws = wb.active
    wb.remove(default_ws)
    _build_sheet_project(wb, schema)
    _build_sheet_paths(wb)
    _build_sheet_doors(wb)
    _build_sheet_readme(wb)
    # Make Sheet 1 the one Excel opens to.
    wb.active = wb.sheetnames.index("Project & Parameters")
    return wb


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate the blank inputs/DiagComm.xlsx template "
        "from assets/DiagComm_schema.json.",
    )
    parser.add_argument(
        "--output", "-o",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Where to write the xlsx. Default: {DEFAULT_OUTPUT.relative_to(SKILL_ROOT).as_posix()}",
    )
    parser.add_argument(
        "--schema",
        type=Path,
        default=SCHEMA_PATH,
        help="Schema file. Default: assets/DiagComm_schema.json",
    )
    args = parser.parse_args(argv)

    if not args.schema.exists():
        sys.stderr.write(
            f"[build_inputs_template] schema not found: {args.schema}\n"
        )
        return 2

    schema = json.loads(args.schema.read_text(encoding="utf-8"))
    wb = build_workbook(schema)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    rel = args.output.relative_to(SKILL_ROOT) if args.output.is_relative_to(SKILL_ROOT) else args.output
    print(f"[build_inputs_template] wrote {rel.as_posix() if hasattr(rel, 'as_posix') else rel}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
