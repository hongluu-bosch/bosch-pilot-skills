"""Tests for ``pipeline.py --init-project``.

Exercises the v2.0.0 workspace bootstrap: scaffolding directories,
copying the blank Excel template, idempotency, ``--force`` overwrite,
and the safety guard that refuses to scaffold inside the skill
checkout.

The fixtures in ``conftest.py`` deliberately monkeypatch
WORKSPACE_ROOT for the regular pipeline commands. ``--init-project``
needs a *real* (un-mocked) cwd-based workspace path to verify, so
these tests use ``monkeypatch.chdir`` to redirect cwd into a tmp
directory and then call ``pipeline.cmd_init_project`` directly --
bypassing the conftest fixture entirely.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = REPO_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import pipeline  # noqa: E402
import runtime   # noqa: E402


def test_init_project_scaffolds_workspace(tmp_path, capsys, monkeypatch):
    """Default invocation with no args targets cwd; creates the four
    workspace subdirs + drops the blank Excel template + writes a
    workspace-local .gitignore."""
    project = tmp_path / "fake_project"
    project.mkdir()
    monkeypatch.chdir(project)

    rc = pipeline.cmd_init_project([])
    assert rc == 0

    workspace = project / ".DCOM_AI" / "DiagComm_Toolkit_PRJ"
    assert workspace.is_dir()
    for sub in ("inputs", "outputs", "state", ".cache"):
        assert (workspace / sub).is_dir(), f"missing {sub}/"

    xlsx = workspace / "inputs" / "DiagComm.xlsx"
    assert xlsx.is_file(), "blank template was not copied"
    # The copied xlsx is *structurally* the template (same sheets, same
    # required-cell layout) but its README sheet has every ``<skill>``
    # placeholder rewritten to the resolved skill path -- so it is NOT
    # byte-identical to the bundled template anymore. Verify the
    # structural contract + the localization side-effect explicitly.
    import openpyxl  # type: ignore[import-not-found]
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        assert {"Project & Parameters", "Paths & Options",
                "DOORS Upload", "README"}.issubset(set(wb.sheetnames))
        readme_text = "\n".join(
            str(c.value) for row in wb["README"].iter_rows()
            for c in row if isinstance(c.value, str))
    finally:
        wb.close()
    assert "<skill>" not in readme_text, \
        "README sheet still contains the <skill> placeholder"
    assert runtime.SKILL_ROOT.as_posix() in readme_text, \
        "README sheet did not get the resolved skill path"

    gi = workspace / ".gitignore"
    assert gi.is_file()
    text = gi.read_text(encoding="utf-8")
    assert ".cache/" in text
    assert "outputs/" in text

    out = capsys.readouterr().out
    assert "Workspace scaffolded." in out
    assert "DiagComm.xlsx" in out


def test_init_project_idempotent(tmp_path, capsys, monkeypatch):
    """Re-running --init-project on an already-initialised workspace
    must NOT overwrite the user's filled Excel and must report
    'Workspace already initialised.'."""
    project = tmp_path / "fake_project"
    project.mkdir()
    monkeypatch.chdir(project)

    pipeline.cmd_init_project([])
    capsys.readouterr()  # discard first-run output

    xlsx = project / ".DCOM_AI" / "DiagComm_Toolkit_PRJ" / "inputs" / "DiagComm.xlsx"
    user_payload = b"USER_FILLED_CONTENT"
    xlsx.write_bytes(user_payload)

    rc = pipeline.cmd_init_project([])
    assert rc == 0

    out = capsys.readouterr().out
    assert "[skip ] inputs/DiagComm.xlsx" in out
    assert "Workspace already initialised." in out
    assert xlsx.read_bytes() == user_payload, \
        "second run overwrote the user's filled template"


def test_init_project_force_overwrites_template(tmp_path, capsys, monkeypatch):
    project = tmp_path / "fake_project"
    project.mkdir()
    monkeypatch.chdir(project)

    pipeline.cmd_init_project([])
    capsys.readouterr()
    xlsx = project / ".DCOM_AI" / "DiagComm_Toolkit_PRJ" / "inputs" / "DiagComm.xlsx"
    xlsx.write_bytes(b"USER_FILLED_CONTENT")

    rc = pipeline.cmd_init_project(["--force"])
    assert rc == 0

    # After --force the workbook is the bundled template again (sheets +
    # required-cell layout intact) and the README sheet has been
    # localized -- same contract as the fresh-scaffold test above. The
    # earlier ``USER_FILLED_CONTENT`` payload must be gone.
    assert xlsx.read_bytes() != b"USER_FILLED_CONTENT"
    import openpyxl  # type: ignore[import-not-found]
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        assert {"Project & Parameters", "Paths & Options",
                "DOORS Upload", "README"}.issubset(set(wb.sheetnames))
        readme_text = "\n".join(
            str(c.value) for row in wb["README"].iter_rows()
            for c in row if isinstance(c.value, str))
    finally:
        wb.close()
    assert "<skill>" not in readme_text
    assert runtime.SKILL_ROOT.as_posix() in readme_text

    out = capsys.readouterr().out
    assert "DiagComm.xlsx" in out


def test_init_project_explicit_root_arg(tmp_path, capsys, monkeypatch):
    """Positional ``project_root`` argument overrides cwd. Useful for
    CI / scripted setup where cwd is some unrelated build dir."""
    project = tmp_path / "alt_project"
    project.mkdir()
    monkeypatch.chdir(tmp_path)  # cwd is *not* the project root

    rc = pipeline.cmd_init_project([str(project)])
    assert rc == 0

    workspace = project / ".DCOM_AI" / "DiagComm_Toolkit_PRJ"
    assert workspace.is_dir()
    assert (workspace / "inputs" / "DiagComm.xlsx").is_file()


def test_init_project_refuses_inside_skill_checkout(tmp_path, capsys, monkeypatch):
    """Refuse to scaffold a workspace whose path is inside the skill
    checkout. Catches the 'oops, forgot to cd to my project root' case
    instead of silently writing project data into the user-level skill."""
    target_root = runtime.SKILL_ROOT / "_pretend_project_inside_skill"

    with pytest.raises(SystemExit) as exc_info:
        pipeline.cmd_init_project([str(target_root)])

    msg = str(exc_info.value)
    assert "refusing to scaffold" in msg
    assert "skill" in msg
    assert "project root" in msg
    # Sanity: nothing was written.
    assert not (target_root / ".DCOM_AI").exists()


def test_init_project_via_main(tmp_path, capsys, monkeypatch):
    """``main()`` short-circuits on ``--init-project`` and routes to
    cmd_init_project before argparse sees the subcommand tree."""
    project = tmp_path / "fake_project"
    project.mkdir()
    monkeypatch.chdir(project)

    rc = pipeline.main(["--init-project"])
    assert rc == 0
    assert (project / ".DCOM_AI" / "DiagComm_Toolkit_PRJ" / "inputs" /
            "DiagComm.xlsx").is_file()

    rc = pipeline.main(["init-project", "--force"])
    assert rc == 0  # also accepts the bare-word form
