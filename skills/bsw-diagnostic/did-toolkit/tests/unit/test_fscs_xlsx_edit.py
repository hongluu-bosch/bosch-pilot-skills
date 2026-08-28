"""Tests for the XLSX-only FSCS edit surface.

The round-trip contract is the same shape it was under the legacy CSV
projection (operator-visible columns only, hidden model preserved on
import) — but the on-disk format is now ``.xlsx`` with data-validation
drop-downs, frozen header row, and autofilter. These tests pin both
the behaviour and the UX scaffolding so a refactor that quietly drops
a drop-down validation or breaks the freeze pane fails loudly.
"""

from __future__ import annotations

import openpyxl

from fscs.xlsx_edit import (
    VALIDATIONS,
    XLSX_COLUMNS,
    export_fscs_xlsx,
    import_fscs_xlsx,
)
from fscs.renderer import render_fscs_json
from fscs.schema import (
    DIDFscsEntry,
    FSCSDocument,
    FSCSFreeText,
    FSCSSecurityLevel,
    FSCSServiceAccess,
    FSCSSubField,
    FSCSValueRangeEnum,
)


def _doc() -> FSCSDocument:
    return FSCSDocument(
        dids=[
            DIDFscsEntry(
                did_hex="F190",
                did_name="BaselineCounter",
                did_name_zh="基线计数器",
                data_type="enum",
                storage_position="EEPROM",
                size_bytes=1,
                rw_state="R",
                nvm_item="NVM_ID_DCOM_BaselineCounter",
                service_22=FSCSServiceAccess(
                    supported=True,
                    used=True,
                    sessions=["defaultSession"],
                    security_levels=[FSCSSecurityLevel(level="L0")],
                    behavior="Read from NVM item: NVM_ID_DCOM_BaselineCounter",
                ),
                service_2e=FSCSServiceAccess(
                    supported=False,
                    used=True,
                    sessions=[],
                    security_levels=[],
                    behavior="",
                ),
                sub_fields=[
                    FSCSSubField(
                        byte_range="0",
                        name_en="counter",
                        range_min="0",
                        range_max="255",
                        enum_mapping="0=zero\n1=one",
                    )
                ],
                value_range=FSCSValueRangeEnum(values=["0", "1"]),
                free_text=FSCSFreeText(description_read="custom read text"),
            )
        ]
    )


def _read_workbook(path):
    """Open ``path`` with openpyxl and return ``(wb, sheet)``.

    Loaded with ``data_only=True`` (so cached cell values, not
    formulae, come back as strings); the read-only flag is OFF so
    tests can inspect ``data_validations`` (read-only mode skips that
    metadata for memory reasons).
    """
    wb = openpyxl.load_workbook(filename=str(path), data_only=True)
    return wb, wb.worksheets[0]


def test_xlsx_export_writes_header_and_row(tmp_path):
    json_path = tmp_path / "fscs.json"
    xlsx_path = tmp_path / "fscs_edit.xlsx"
    json_path.write_text(render_fscs_json(_doc()), encoding="utf-8")

    assert export_fscs_xlsx(json_path, xlsx_path) == 1

    wb, sheet = _read_workbook(xlsx_path)
    try:
        header = [cell.value for cell in sheet[1]]
        assert header == XLSX_COLUMNS

        row = [cell.value for cell in sheet[2]]
        row_map = dict(zip(header, row))
        assert row_map["did_hex"] == "0xF190"
        assert row_map["used_flag"] == "TRUE"
        assert row_map["product_type"] == "Common"
        assert row_map["service_22_behavior"].startswith("Read from NVM item:")
        assert row_map["service_2e_behavior"] == ""
        assert "value_range_json" not in row_map
        assert "sub_fields_json" not in row_map
    finally:
        wb.close()


def test_xlsx_export_attaches_dropdown_for_every_constrained_column(tmp_path):
    json_path = tmp_path / "fscs.json"
    xlsx_path = tmp_path / "fscs_edit.xlsx"
    json_path.write_text(render_fscs_json(_doc()), encoding="utf-8")
    export_fscs_xlsx(json_path, xlsx_path)

    wb, sheet = _read_workbook(xlsx_path)
    try:
        # Collect every column index touched by a list-type data
        # validation and the canonical value list each one carries.
        # openpyxl exposes the formula1 source as a comma-joined
        # quoted string (e.g. '"TRUE,FALSE"').
        seen: dict[int, tuple[str, ...]] = {}
        for dv in sheet.data_validations.dataValidation:
            if dv.type != "list":
                continue
            raw = (dv.formula1 or "").strip().strip('"')
            allowed = tuple(part.strip() for part in raw.split(","))
            for cell_range in dv.sqref.ranges:
                # cell_range.min_col is 1-indexed; XLSX_COLUMNS is 0-indexed
                col_idx = cell_range.min_col - 1
                seen[col_idx] = allowed

        for column, expected in VALIDATIONS.items():
            col_idx = XLSX_COLUMNS.index(column)
            assert col_idx in seen, f"missing drop-down on column {column!r}"
            assert seen[col_idx] == expected, (
                f"drop-down for {column!r}: expected {expected}, got {seen[col_idx]}"
            )
    finally:
        wb.close()


def test_xlsx_export_freezes_header_and_attaches_autofilter(tmp_path):
    json_path = tmp_path / "fscs.json"
    xlsx_path = tmp_path / "fscs_edit.xlsx"
    json_path.write_text(render_fscs_json(_doc()), encoding="utf-8")
    export_fscs_xlsx(json_path, xlsx_path)

    wb, sheet = _read_workbook(xlsx_path)
    try:
        # Freeze panes set to A2 means row 1 (the header) is frozen.
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
    finally:
        wb.close()


def test_xlsx_import_round_trips_visible_edits_and_preserves_hidden_model(tmp_path):
    json_path = tmp_path / "fscs.json"
    xlsx_path = tmp_path / "fscs_edit.xlsx"
    json_path.write_text(render_fscs_json(_doc()), encoding="utf-8")
    export_fscs_xlsx(json_path, xlsx_path)

    # Mutate a few operator-facing cells in place, then save the
    # workbook back via openpyxl (the simulation of what Excel would
    # write on disk after an operator edits).
    wb = openpyxl.load_workbook(filename=str(xlsx_path))
    try:
        sheet = wb.worksheets[0]
        header = [cell.value for cell in sheet[1]]
        col = {h: i + 1 for i, h in enumerate(header)}
        row = 2
        sheet.cell(row=row, column=col["used_flag"]).value = "FALSE"
        sheet.cell(row=row, column=col["did_name"]).value = "UpdatedCounter"
        sheet.cell(row=row, column=col["service_2e_support"]).value = "TRUE"
        sheet.cell(row=row, column=col["service_22_behavior"]).value = "custom read behavior"
        sheet.cell(row=row, column=col["service_2e_behavior"]).value = "custom write behavior"
        wb.save(str(xlsx_path))
    finally:
        wb.close()

    updated = import_fscs_xlsx(json_path, xlsx_path)
    did = updated.dids[0]

    assert did.did_name == "UpdatedCounter"
    assert did.service_22.used is False
    assert did.service_2e.used is False
    assert did.service_2e.supported is True
    assert did.service_22.behavior == "custom read behavior"
    assert did.service_2e.behavior == "custom write behavior"
    # Hidden / non-operator fields survive the round-trip untouched.
    assert did.value_range.values == ["0", "1"]
    assert did.sub_fields[0].enum_mapping == "0=zero\n1=one"
    assert did.free_text.description_read == "custom read text"
